import streamlit as st
import datetime
import holidays
import pandas as pd
from docxtpl import DocxTemplate
from bahttext import bahttext
import os
import google.generativeai as genai
from PIL import Image
import json
from docx.shared import Cm, Pt

# ==========================================
# 1. ตั้งค่าหน้าจอ UI & CSS
# ==========================================
st.set_page_config(page_title="ระบบทำเอกสารพัสดุ 9 ขั้นตอน", page_icon="📅", layout="wide")

st.markdown("""
    <style>
    .stApp { background-color: #FFFFFF; }
    .stTextInput input, .stSelectbox select, .stTextArea textarea, .stDateInput input {
        background-color: #FFFFFF !important; 
        color: #333333 !important;           
        border-radius: 10px !important;
        border: 1px solid #E0E0E0 !important;
        padding: 10px !important;
    }
    .stDataEditor {
        background-color: #FFFFFF !important;
        border-radius: 10px !important;
    }
    .header-zone {
        padding: 15px 25px;
        border-radius: 12px;
        margin-bottom: 25px;
        color: white !important;
        font-weight: 600;
        box-shadow: 0 4px 15px rgba(0,0,0,0.08);
    }
    .header-green { background: linear-gradient(135deg, #1D976C 0%, #93F9B9 100%); }
    .header-blue { background: linear-gradient(135deg, #2193b0 0%, #6dd5ed 100%); }
    .header-orange { background: linear-gradient(135deg, #FF8C00 0%, #FFD700 100%); }
    .header-zone h3 {
        margin: 0 !important;
        color: white !important;
        font-size: 1.25rem !important;
        letter-spacing: 0.5px;
    }
    div.row-widget.stRadio > div { flex-direction:row; }
    .stButton>button {
        border-radius: 12px;
        padding: 15px 30px;
        font-weight: 700;
        transition: all 0.3s;
    }
    .stButton>button:hover {
        transform: translateY(-2px);
        box-shadow: 0 5px 15px rgba(0,0,0,0.1);
    }
    </style>
""", unsafe_allow_html=True)

if 'data' not in st.session_state:
    st.session_state.data = {"doc_type": "จัดซื้อ", "shop_name": "", "items": []}

st.write(f'<p style="font-size:36px; font-weight:bold; color:#2C3E50; margin-bottom:5px;">📅 ระบบจัดทำเอกสารพัสดุ</p>', unsafe_allow_html=True)
st.write(f'<p style="color:#7F8C8D; font-size:18px;">สะดวก รวดเร็ว และเป็นระเบียบตามระเบียบพัสดุ</p>', unsafe_allow_html=True)

st.sidebar.header("⚙️ ตั้งค่าระบบ AI")
api_key = st.sidebar.text_input("🔑 ใส่ Gemini API Key:", type="password")
st.markdown("---")

# ==========================================
# 1. นำเข้าข้อมูล (Input & AI Predictor)
# ==========================================
st.markdown('<div class="header-zone header-blue"><h3>📸  นำเข้าข้อมูลพัสดุ (AI แกะลายมือ)</h3></div>', unsafe_allow_html=True)

col_cam, col_file = st.columns(2)
with col_cam:
    camera_photo = st.camera_input("ถ่ายรูปบิล")
with col_file:
    uploaded_photo = st.file_uploader("อัปโหลดรูปภาพ", type=['png', 'jpg', 'jpeg'])

image_to_process = camera_photo if camera_photo else uploaded_photo
# 🌟 1. ฐานข้อมูลรายชื่อบุคลากร (คุณครูสามารถแก้ไขชื่อจริง-นามสกุลจริงได้ตรงนี้เลยครับ)
TEACHER_LIST = [
    "รอข้อมูลAI",
    "นายเชาว์  แดขุนทด",
    "นางสาวกนกภรณ์ พัฒนาศูนย์",
    "นางสาวชมชนก ชัยปัญญา",
    "นายณัฐพล กันพล ",
    "นางธัญญารัตน์ หมื่นภักดี",
    "นายธีระพงษ์ คำพระธิก",
    "นางนงนุช บัวงาม"
    "นางสาวรัชดาวรรณ นาเมฆ",
    "นางรัชนก คำพระธิก ",
    "นายวชิระวิชญ์ หินซุยอัครภา"
    "นางสาวศลิษา พลเขต",
    "นางสาวสายฝน เจริญ",
    "นางสาวสุภาวดี เหลาบับภา"
    "นางสาวโสภา กุลากุล",
    "นางอมรรัตน์ ทิพย์พรมมา",
    "นางสาวอุดมลักษณ์ ประพงษ์"
]
if image_to_process is not None and st.button("🤖 ให้ AI แกะข้อมูลจากรูปภาพ", type="primary", use_container_width=True):
    if not api_key:
        st.error("⚠️ กรุณาใส่ API Key ที่เมนูซ้ายมือก่อนครับ")
    else:
        try:
            with st.spinner("AI กำลังอ่านและวิเคราะห์ประเภทเอกสาร... 🪄"):
                genai.configure(api_key=api_key)
                model = genai.GenerativeModel('gemini-3.1-flash-image-preview')
                img = Image.open(image_to_process)
                
                # แปลงรายชื่อครูเป็นข้อความยาวๆ เพื่อส่งให้ AI อ่าน
                teachers_str = ", ".join(TEACHER_LIST)
                
                prompt = f"""
                จงอ่านข้อมูลจากภาพแบบฟอร์มขอซื้อ/ขอจ้างนี้ และสรุปเป็นรูปแบบ JSON เท่านั้น ห้ามมีข้อความอื่นปน 
                
                🚨 กฎสำคัญเรื่องรายชื่อกรรมการ: 
                นี่คือรายชื่อบุคลากรในโรงเรียนที่ถูกต้อง: [{teachers_str}]
                หากในเอกสารมีการระบุชื่อกรรมการ ให้คุณหาชื่อที่ "ใกล้เคียงที่สุด" จากโพยด้านบน แล้วนำ "ชื่อเต็มพร้อมตำแหน่งจากโพย" มาใส่ในช่อง inspector_1, 2, 3 ให้ถูกต้องเป๊ะๆ (ถ้าไม่มีชื่อให้ใส่คำว่า "ไม่ระบุ")
                
                โดยใช้โครงสร้างดังนี้:
                {{
                    "doc_type": "ให้วิเคราะห์จากรายการสินค้าว่าเป็น 'จัดซื้อ' หรือ 'จัดจ้าง'",
                    "item_title": "ชื่อรายการใหญ่ที่จะขอซื้อ/จ้าง",
                    "shop_name": "ชื่อร้านค้า/บริษัท",
                    "seller_name": "ชื่อตัวบุคคลผู้ขาย/ผู้รับจ้าง (ถ้ามี)",
                    "vendor_no": "ที่อยู่ เลขที่",
                    "vendor_moo": "หมู่ที่",
                    "vendor_subdistrict": "ตำบล",
                    "vendor_district": "อำเภอ",
                    "vendor_province": "จังหวัด",
                    "vendor_phone": "เบอร์โทร",
                    "vendor_tax_id": "เลขประจำตัวผู้เสียภาษี",
                    "bank_name": "ชื่อธนาคาร",
                    "bank_branch": "สาขาธนาคาร",
                    "bank_account": "เลขที่บัญชี",
                    "project_name": "ชื่อโครงการ",
                    "activity_name": "ชื่อกิจกรรม",
                    "department": "กลุ่มงาน/ฝ่าย",
                    "purchase_reason": "เหตุผลที่ขอ",
                    "inspector_1": "ชื่อประธานกรรมการ (ดึงจากโพยเท่านั้น)",
                    "inspector_2": "ชื่อกรรมการคนที่ 2 (ดึงจากโพยเท่านั้น)",
                    "inspector_3": "ชื่อกรรมการคนที่ 3 (ดึงจากโพยเท่านั้น)",
                    "items": [
                        {{"name": "ชื่อสินค้าหรือชื่องาน", "qty": 1, "unit": "หน่วย", "price": 100, "total": 100}}
                    ]
                }}
                """
                response = model.generate_content([prompt, img])
                result_text = response.text.replace('```json', '').replace('```', '').strip()
                extracted_data = json.loads(result_text)
                
                st.session_state.data.update(extracted_data)
                st.success(f"✨ AI แกะข้อมูลสำเร็จ! (ระบบวิเคราะห์ว่าเป็น: การ{extracted_data.get('doc_type', 'จัดซื้อ')})")
        except Exception as e:
            st.error(f"❌ เกิดข้อผิดพลาดในการประมวลผลรูปภาพ: {str(e)}")

st.markdown("---")

# ==========================================
# ฟังก์ชันจัดการวันที่
# ==========================================
def get_past_working_day(start_date, days_back):
    th_holidays = holidays.Thailand()
    current_date = start_date
    days_counted = 0
    while days_counted < days_back:
        current_date -= datetime.timedelta(days=1)
        if current_date.weekday() < 5 and current_date not in th_holidays:
            days_counted += 1
    return current_date

def format_thai_date(date_obj):
    thai_months = ["", "มกราคม", "กุมภาพันธ์", "มีนาคม", "เมษายน", "พฤษภาคม", "มิถุนายน", 
                   "กรกฎาคม", "สิงหาคม", "กันยายน", "ตุลาคม", "พฤศจิกายน", "ธันวาคม"]
    return f"{date_obj.day} {thai_months[date_obj.month]} {date_obj.year + 543}"

# ==========================================
# 2. ส่วนการตรวจสอบและแก้ไข
# ==========================================
col_left, col_right = st.columns([1, 2], gap="large")

with col_left:
    st.markdown('<div class="header-zone header-green"><h3>🗓️ 1. กำหนดวันเบิกเงิน</h3></div>', unsafe_allow_html=True)
    base_pay_date = st.date_input("วันที่บันทึกรายงานผล / ขอเบิกเงิน (วันสุดท้าย)", datetime.date.today())
    
    st.write("")
    
    st.markdown('<div class="header-zone header-blue"><h3>⏳ 2. ไทม์ไลน์ 9 ขั้นตอน</h3></div>', unsafe_allow_html=True)
    
    d_ins = get_past_working_day(base_pay_date, 2)
    d_dev = get_past_working_day(d_ins, 2)
    d_po  = get_past_working_day(d_dev, 2)
    d_ann = get_past_working_day(d_po, 2)
    d_app = get_past_working_day(d_ann, 2)
    d_ord = get_past_working_day(d_app, 2)
    d_quo = get_past_working_day(d_ord, 2)
    d_req = get_past_working_day(d_quo, 2)

    req_d = st.date_input("1. วันที่รายงานขอซื้อ/ขอจ้าง", d_req)
    quo_d = st.date_input("2. วันที่ใบเสนอราคา", d_quo)
    ord_d = st.date_input("3. วันที่คำสั่งแต่งตั้ง", d_ord)
    app_d = st.date_input("4. วันที่พิจารณาเห็นชอบ", d_app)
    ann_d = st.date_input("5. วันที่ประกาศผู้ชนะ", d_ann)
    po_d  = st.date_input("6. วันที่ใบสั่งซื้อ/สั่งจ้าง", d_po)
    dev_d = st.date_input("7. วันที่ส่งมอบงาน/ของ", d_dev)
    ins_d = st.date_input("8. วันที่ตรวจรับพัสดุ", d_ins)
    pay_d = st.date_input("9. วันที่ขอเบิกเงิน", base_pay_date)

with col_right:
    st.markdown('<div class="header-zone header-orange"><h3>📝 3. ข้อมูลผู้ขายและโครงการ</h3></div>', unsafe_allow_html=True)
    
    ai_doc_type = st.session_state.data.get('doc_type', 'จัดซื้อ')
    type_index = 0 if ai_doc_type == "จัดซื้อ" else 1

    st.markdown("#### 🛒 ประเภทการจัดหา")
    doc_type = st.radio("เลือกประเภท:", ["จัดซื้อ", "จัดจ้าง"], index=type_index, horizontal=True, label_visibility="collapsed")
    
    if doc_type == "จัดซื้อ":
        w_buy, w_vendor, w_po, w_delivery = "ซื้อ", "ผู้ขาย", "ใบสั่งซื้อ", "ใบส่งของ"
    else:
        w_buy, w_vendor, w_po, w_delivery = "จ้าง", "ผู้รับจ้าง", "ใบสั่งจ้าง", "ใบส่งมอบงาน"

    st.markdown(f"#### 🏢 ข้อมูลผู้ขาย/ผู้รับจ้าง ({w_vendor})")
    col_v1, col_v2 = st.columns(2)
    with col_v1:
        f_shop = st.text_input(f"ชื่อร้านค้า/บริษัท", st.session_state.data.get('shop_name', ''))
    with col_v2:
        f_seller = st.text_input(f"ชื่อตัวบุคคลผู้ขาย/ผู้รับจ้าง", st.session_state.data.get('seller_name', ''))
    
    # 🌟 เอาที่อยู่แยกส่วนกลับมาให้ครบถ้วนครับ!
    st.markdown("**📍 ที่อยู่และข้อมูลติดต่อ**")
    addr1, addr2, addr3 = st.columns(3)
    with addr1:
        f_no = st.text_input("เลขที่", st.session_state.data.get('vendor_no', ''))
        f_subdistrict = st.text_input("ตำบล", st.session_state.data.get('vendor_subdistrict', ''))
    with addr2:
        f_moo = st.text_input("หมู่ที่", st.session_state.data.get('vendor_moo', ''))
        f_district = st.text_input("อำเภอ", st.session_state.data.get('vendor_district', ''))
    with addr3:
        f_province = st.text_input("จังหวัด", st.session_state.data.get('vendor_province', ''))
        f_phone = st.text_input("เบอร์โทรศัพท์", st.session_state.data.get('vendor_phone', ''))

    f_tax = st.text_input("เลขผู้เสียภาษี", st.session_state.data.get('vendor_tax_id', ''))

    st.markdown("#### 🏦 ข้อมูลบัญชีธนาคาร")
    cb1, cb2, cb3 = st.columns(3)
    with cb1:
        f_bank_name = st.text_input("ชื่อธนาคาร", st.session_state.data.get('bank_name', ''))
    with cb2:
        f_bank_branch = st.text_input("สาขา", st.session_state.data.get('bank_branch', ''))
    with cb3:
        f_bank_account = st.text_input("เลขบัญชี", st.session_state.data.get('bank_account', ''))

    st.markdown("#### 📦 รายการที่จะขอซื้อ/จ้าง")
    f_item_title = st.text_input("ชื่อรายการใหญ่ (เช่น วัสดุสำนักงาน, จ้างทำป้าย)", st.session_state.data.get('item_title', ''))
    
    items_data = st.session_state.data.get('items', [{"name": "", "qty": 1, "unit": "", "price": 0, "total": 0}])
    edited_df = st.data_editor(pd.DataFrame(items_data), num_rows="dynamic", use_container_width=True)
    
    st.markdown("#### 💰 การคำนวณภาษีและยอดเงิน")
    has_vat = st.checkbox("ร้านค้านี้จดทะเบียนภาษีมูลค่าเพิ่ม (ถอด VAT 7% ในตัว)", value=False)
    
    try:
        raw_total = edited_df['total'].sum()
    except:
        raw_total = 0.0
        
    if has_vat:
        before_vat = raw_total / 1.07
        vat_amount = raw_total - before_vat
    else:
        before_vat = raw_total
        vat_amount = 0.0
        
    st.info(f"ราคาก่อนภาษี: {before_vat:,.2f} | ภาษี (7%): {vat_amount:,.2f} | ยอดสุทธิ: {raw_total:,.2f} บาท ({bahttext(raw_total)})")

    st.markdown("#### 📋 ข้อมูลโครงการและกรรมการ")
    c3, c4 = st.columns(2)
    with c3:
        f_project = st.text_input("ชื่อโครงการ", st.session_state.data.get('project_name', ''))
        f_activity = st.text_input("ชื่อกิจกรรม", st.session_state.data.get('activity_name', ''))
    with c4:
        f_dept = st.text_input("กลุ่มงาน/ฝ่าย", st.session_state.data.get('department', ''))
        f_delivery_days = st.text_input("กำหนดส่งมอบ (เช่น 7 วัน, 15 วัน)", "7 วัน")

    f_reason = st.text_area("เหตุผลความจำเป็นที่ขอซื้อ/จ้าง", st.session_state.data.get('purchase_reason', ''))

    st.markdown("**รายชื่อคณะกรรมการตรวจรับพัสดุ**")
    
    # ฟังก์ชันช่วยค้นหาว่าชื่อที่ AI ส่งมา ตรงกับลำดับที่เท่าไหร่ในโพย
    def get_teacher_index(ai_name):
        return TEACHER_LIST.index(ai_name) if ai_name in TEACHER_LIST else 0

    c5, c6, c7 = st.columns(3)
    with c5:
        # ดึงชื่อที่ AI เดามา (ถ้าไม่มีค่าให้เป็น "ไม่ระบุ")
        ai_insp_1 = st.session_state.data.get('inspector_1', 'ไม่ระบุ')
        # สร้าง Dropdown เมนู
        f_inspector_1 = st.selectbox("ประธานกรรมการ", TEACHER_LIST, index=get_teacher_index(ai_insp_1))
    with c6:
        ai_insp_2 = st.session_state.data.get('inspector_2', 'ไม่ระบุ')
        f_inspector_2 = st.selectbox("กรรมการคนที่ 2", TEACHER_LIST, index=get_teacher_index(ai_insp_2))
    with c7:
        ai_insp_3 = st.session_state.data.get('inspector_3', 'ไม่ระบุ')
        f_inspector_3 = st.selectbox("กรรมการคนที่ 3", TEACHER_LIST, index=get_teacher_index(ai_insp_3))

# ==========================================
# 3. ส่วนการสร้างเอกสาร (Output)
# ==========================================
st.markdown("---")
if st.button("🖨️ สร้างเอกสารพัสดุครบชุด (10 หน้า)", type="primary", use_container_width=True):
    try:
        items_list = edited_df.to_dict('records')
        doc = DocxTemplate("templat patsadu.docx")
        
        final_data = {
            "buy_or_hire": w_buy,
            "vendor_type": w_vendor,
            "po_name": w_po,
            "delivery_name": w_delivery,
            "total_item_count": str(len(items_list)),
            
            # 🌟 อัปเดต: ส่งข้อมูลที่อยู่แบบแยกส่วนกลับไปที่ Word ครบถ้วน
            "vendor_name": f_shop, 
            "seller_name": f_seller,
            "vendor_no": f_no, 
            "vendor_moo": f_moo,          
            "vendor_subdistrict": f_subdistrict,  
            "vendor_district": f_district,     
            "vendor_province": f_province,     
            "vendor_phone": f_phone, 
            "vendor_tax_id": f_tax, 
            
            "bank_name": f_bank_name,
            "bank_branch": f_bank_branch,
            "bank_account": f_bank_account,    
            
            "item_title": f_item_title,
            "project_name": f_project, 
            "activity_name": f_activity,
            "department": f_dept,
            "purchase_reason": f_reason,
            "delivery_days": f_delivery_days,
            
            "inspector_1": f_inspector_1,
            "inspector_2": f_inspector_2,
            "inspector_3": f_inspector_3,
            
            "price_before_vat": f"{before_vat:,.2f}",
            "vat_amount": f"{vat_amount:,.2f}",
            "budget_amount": f"{raw_total:,.2f}", 
            "budget_amount_text": bahttext(raw_total),
            "withholding_tax_amount": "-", 
            "fine_amount": "-", 
            "net_payable_amount": f"{raw_total:,.2f}", 
            "net_payable_amount_text": bahttext(raw_total),
            
            "request_date": format_thai_date(req_d), 
            "quote_date": format_thai_date(quo_d),
            "order_date": format_thai_date(ord_d), 
            "approval_date": format_thai_date(app_d),
            "announce_date": format_thai_date(ann_d), 
            "po_date": format_thai_date(po_d),
            "delivery_date": format_thai_date(dev_d), 
            "inspect_date": format_thai_date(ins_d),
            "pay_date": format_thai_date(pay_d)
        }
        
        doc.render(final_data)
        
        for table in doc.docx.tables:
            if len(table.rows) > 0 and "รายการ" in table.cell(0, 1).text:
                for i, item in enumerate(items_list, 1):
                    row_cells = table.add_row().cells
                    if len(row_cells) >= 6:
                        values = [
                            str(i),
                            str(item.get('name', '')),
                            str(item.get('qty', '')),
                            str(item.get('unit', '')),
                            str(item.get('price', '')),
                            str(item.get('total', ''))
                        ]
                        
                        for idx, val in enumerate(values):
                            row_cells[idx].text = val
                            for paragraph in row_cells[idx].paragraphs:
                                for run in paragraph.runs:
                                    run.font.name = 'TH Sarabun PSK'
                                    run.font.size = Pt(15)

        output_file = f"เอกสารพัสดุ_{f_shop}.docx"
        doc.save(output_file)
        
        st.success(f"✅ สร้างเอกสารประเภท 'การ{doc_type}' สำเร็จ!")
        with open(output_file, "rb") as f:
            st.download_button("📥 ดาวน์โหลดไฟล์ Word", f, file_name=output_file, use_container_width=True)
            
    except Exception as e:
        st.error(f"❌ เกิดข้อผิดพลาด: {str(e)}")
